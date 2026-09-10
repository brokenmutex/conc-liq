"""Render frozen inventory research; install matplotlib and openpyxl in a local venv."""
import collections,csv,hashlib,json,math,pathlib,statistics,sys
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
src=pathlib.Path(sys.argv[1]);out=pathlib.Path(sys.argv[2]);out.mkdir(parents=True,exist_ok=True)
read=lambda n:json.loads((src/n).read_text())
a=read('inventory-study-v2.json');metrics=read('session-metrics.json');ops=read('operations-checkpoint.json');fix=json.loads(pathlib.Path('test/fixtures/inventory-recenter-fork.json').read_text())
unit=lambda v:int(v)/1e6

def write(n,obj): (out/n).write_text(json.dumps(obj,indent=2)+'\n')
def table_csv(name,rows):
 with (out/name).open('w') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]),lineterminator='\n');w.writeheader();w.writerows(rows)
rows=[]
for r in a['rows']:
 row={k:v for k,v in r.items() if k not in ['actions','rejected','terminalExecution']}
 row['rejected']='; '.join(r['rejected']);row['actions']='; '.join(x['kind']+'@'+x['checkpoint'] for x in r['actions'])
 for k in ['net','alpha','newGas','feesAfterAnchor','maxDrawdown','delta']:row[k+'_USDG']=None if row[k] is None else unit(row.pop(k))
 rows.append(row)
table_csv('episodes.csv',rows);table_csv('sessions.csv',a['sessionRows']);table_csv('aggregates.csv',a['aggregates'])
summary={'asOf':a['asOf'],'source':a['source'],'method':a['method'],'campaign':metrics['campaign'],'totals':metrics['totals'],'sessions':a['sessionRows'],'aggregates':a['aggregates'],'profiles':a['profiles']}
write('study-summary.json',summary)
write('operations-checkpoint.json',ops)
write('accounting-confirmation.json',read('accounting-confirmation.json'))
parity=read('recorded-action-parity.json');write('replay-parity.json',{k:v for k,v in parity.items() if k not in ['native','sessions']}|{'sessions':[{k:v for k,v in x.items() if k!='marks'}|{'verifiedMarks':len(x['marks'])} for x in parity['sessions']]})
mon=[json.loads(l) for l in (src/'holding-observations.jsonl').read_text().splitlines()]
write('holding-observation-extract.json',[{k:o[k] for k in ['checkedAt','sessionId','status','action','source','holding','nav','costs','checks']} for o in mon if '2026-09-10T13:00:00'<=o['checkedAt']<='2026-09-10T13:03:00'])
proof=[]
for name in ['exit','recenter','preserve','trim']:
 x=read(name+'-54.json');r=x['result'];proof.append({k:v for k,v in x.items() if k not in ['result','oracleEvidence']}|{'source':r['source'],'balances':r['balances'],'plan':r.get('plan'),'position':r.get('position'),'retainedLiquidity':r.get('remainingLiquidity'),'principalAfter':r.get('principalAfter'),'priceAfter':r.get('priceAfter'),'transactions':[{'action':t['action'],'totalFeeWei':t['estimate']['totalFeeWei']} for t in r['transactions']],'upstream':r['upstream'],'limitations':r['limitations']})
write('fork-evidence.json',proof)
# Own integer v3 principal calculation. Extreme marks are scenarios, not future quotes.
Q96=1<<96
# Recover exact TickMath endpoints using the fixture tick prices from a Node export.
ratios=read('stress-tick-ratios.json');ratio=lambda t:int(ratios[str(t)])
p0=int(fix['seed']['price']);ref0=(1<<192)*10**30//p0**2
c=fix['cases'][0];i=c['inventory'];prior_gas=260684
base={'name':'Retain original LP','cash':int(i['idle0']),'rwa':int(i['idle1']),'lower':i['tickLower'],'upper':i['tickUpper'],'L':int(i['liquidity']),'fee0':int(i['fee0']),'fee1':int(i['fee1']),'gas':prior_gas}
books=[base]
for c in fix['cases']:
 p=c['plan'];books.append({'name':'Balanced recenter' if c['action']=='recenter' else 'Recenter without sale','cash':int(p['mint']['idle0']),'rwa':int(p['mint']['idle1']),'lower':p['range']['tickLower'],'upper':p['range']['tickUpper'],'L':int(p['mint']['liquidity']),'fee0':0,'fee1':0,'gas':prior_gas+int(c['gasQuote'])})
t=fix['trim'];books.append({'name':'Trim 25% and sell','cash':int(t['balances']['after']['quote']),'rwa':int(t['balances']['after']['rwa']),'lower':i['tickLower'],'upper':i['tickUpper'],'L':int(t['remainingLiquidity']),'fee0':0,'fee1':0,'gas':prior_gas+int(t['gasQuote'])})
def balances(b,price):
 lo,hi=ratio(b['lower']),ratio(b['upper']);p=min(max(price,lo),hi)
 return b['cash']+((b['L']<<96)*(hi-p)//hi)//p+b['fee0'],b['rwa']+b['L']*(p-lo)//Q96+b['fee1']
def mark(b,price,reference):
 q,r=balances(b,price);return q+r*reference//10**30-b['gas']
stress=[]
for b in books:
 v0=mark(b,p0,ref0);q,r=balances(b,p0);exp=r*ref0//10**30*1e6/v0
 for loss in [0,1,5,10]:
  price=math.isqrt(p0*p0*100//(100-loss));reference=ref0*(100-loss)//100;v=mark(b,price,reference)
  stress.append({'action':b['name'],'priceFallPercent':loss,'navUSDG':v/1e6,'lossFromAfterActionUSDG':(v0-v)/1e6,'startingNvdaExposurePercent':exp/10000})
table_csv('mechanical-stress.csv',stress)
wb=Workbook();wb.remove(wb.active)
for name,records in [('Episodes',rows),('Sessions',a['sessionRows']),('Aggregates',a['aggregates']),('Mechanical stress',stress),('Action costs',[{'action':x['action'],'sourceBlock':x['source']['block'],'gas_USDG':unit(x['gasQuote']),'transactions':len(x['transactions'])} for x in proof])]:
 ws=wb.create_sheet(name);keys=list(records[0]);ws.append(keys)
 for row in records:ws.append([json.dumps(row.get(k)) if isinstance(row.get(k),(dict,list)) else row.get(k) for k in keys])
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24425C')
 for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(35,max(12,len(str(col[0].value))+2))
ws=wb.create_sheet('Read me');ws.append(['Item','Value']);ws.append(['Scope','Retrospective overlapping scenarios; do not sum as campaign PnL']);ws.append(['Paper as of',a['asOf']]);ws.append(['Market hash',a['source']['marketSha256']]);ws.append(['Costs','Source block scenarios, not observed future gas']);ws.append(['Stress','Integer v3 mark only; no fees, gas changes or future liquidity quotes']);ws.append(['Units','Episodes are USDG; Sessions raw monetary fields are micro-USDG; ppm are parts per million'])
wb.save(out/'inventory-study.xlsx')
plt.rcParams.update({'figure.dpi':160,'font.size':10,'axes.spines.top':False,'axes.spines.right':False})
fig,axes=plt.subplots(1,2,figsize=(11,4.4));groups=['inventory','chain','risk_reference','operator'];labels=['Inventory','Chain','Risk/reference','Operator']
for j,key in enumerate(['fees','gas','pnl']):
 vals=[sum(unit(r[key]) for r in a['sessionRows'] if r['group']==g) for g in groups];axes[0].bar([i+(j-1)*.23 for i in range(4)],vals,.23,label=key.capitalize())
axes[0].set_xticks(range(4),labels,rotation=15);axes[0].set_ylabel('USDG across closed sessions');axes[0].axhline(0,color='#888',lw=.8);axes[0].legend();axes[0].set_title('Recorded paper accounting: 50 sessions')
strategies=['exit_early','recenter60','trim25','hold80'];names=['Exit at 50%','Recenter','Trim 25%','Hold to 80%']
for j,h in enumerate([30,120]):
 vals=[next(x for x in a['aggregates'] if x['key']==f'500000/{h}/source_block_54/1000000/0/{s}')['meanDelta'] for s in strategies];axes[1].bar([i+(j-.5)*.32 for i in range(4)],vals,.32,label=f'{h} minutes (n={19 if h==30 else 17})')
axes[1].set_xticks(range(4),names,rotation=15);axes[1].set_ylabel('Mean USDG difference versus 60% guard');axes[1].axhline(0,color='#888',lw=.8);axes[1].legend();axes[1].set_title('Overlapping scenarios, newer cost profile')
fig.tight_layout();fig.savefig(out/'performance.png');plt.close(fig)
fig,ax=plt.subplots(figsize=(8,4.5))
for b in books:
 points=[r for r in stress if r['action']==b['name']];ax.plot([r['priceFallPercent'] for r in points],[r['lossFromAfterActionUSDG'] for r in points],marker='o',label=b['name'])
ax.set_xlabel('Hypothetical NVDA price fall (%)');ax.set_ylabel('Additional marked loss after action (USDG)');ax.set_title('Case 54: full v3 inventory conversion in a price decline');ax.legend();fig.tight_layout();fig.savefig(out/'inventory-stress.png');plt.close(fig)
manifest={n.name:hashlib.sha256(n.read_bytes()).hexdigest() for n in src.iterdir() if n.is_file() and n.suffix in ['.json','.jsonl'] and n.name!='paper-source.json'};manifest['paper-source.json']=a['source']['paperSha256'];write('source-manifest.json',manifest)
print(json.dumps({'artifacts':str(out),'rows':len(rows),'stress':stress},indent=2))
